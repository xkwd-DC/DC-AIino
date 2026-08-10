"""
Convert 软件设计说明书.md to properly formatted .docx
with headers, footers, page numbers, and Chinese-friendly formatting.
"""
import re
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

def set_cell_shading(cell, color):
    """Set cell background color."""
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), color)
    shading.set(qn('w:val'), 'clear')
    cell._tc.get_or_add_tcPr().append(shading)

def add_page_number(doc):
    """Add page numbers to footer."""
    for section in doc.sections:
        footer = section.footer
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = p.add_run("— ")
        run.font.size = Pt(9)
        run.font.name = 'SimSun'

        fldChar1 = OxmlElement('w:fldChar')
        fldChar1.set(qn('w:fldCharType'), 'begin')
        run._r.append(fldChar1)

        instrText = OxmlElement('w:instrText')
        instrText.set(qn('xml:space'), 'preserve')
        instrText.text = ' PAGE '
        run._r.append(instrText)

        fldChar2 = OxmlElement('w:fldChar')
        fldChar2.set(qn('w:fldCharType'), 'end')
        run._r.append(fldChar2)

        run2 = p.add_run(" —")
        run2.font.size = Pt(9)
        run2.font.name = 'SimSun'

def add_header(doc, software_name):
    """Add header with software name and copyright holder."""
    for section in doc.sections:
        header = section.header
        header.is_linked_to_previous = False
        p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"{software_name} V1.0 — 软件设计说明书 — 第一著作权人:潘妙齐")
        run.font.size = Pt(8)
        run.font.name = 'SimSun'
        run.font.color.rgb = RGBColor(128, 128, 128)

def add_paragraph(doc, text, font_size=11, font_name='SimSun', alignment=None, space_after=6):
    """Add a paragraph with specified formatting."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(font_size)
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    if alignment is not None:
        p.alignment = alignment
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    return p

def add_heading_styled(doc, text, level=1):
    """Add a heading with Chinese-friendly styling."""
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = 'SimHei'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimHei')
    return h

def add_table_from_markdown(doc, md_table_text):
    """Parse a markdown table and add it to the document."""
    lines = [l.strip() for l in md_table_text.strip().split('\n') if l.strip() and not l.strip().startswith('|--')]
    if not lines:
        return

    rows = []
    for line in lines:
        cells = [c.strip() for c in line.split('|') if c.strip()]
        if cells:
            rows.append(cells)

    if not rows:
        return

    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = 'Table Grid'

    for i, row_data in enumerate(rows):
        for j, cell_text in enumerate(row_data):
            cell = table.cell(i, j)
            cell.text = cell_text
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                for run in p.runs:
                    run.font.size = Pt(9)
                    run.font.name = 'SimSun'
                    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')
                    if i == 0:
                        run.bold = True

            if i == 0:
                set_cell_shading(cell, 'D9E2F3')

    doc.add_paragraph()

def process_markdown_to_docx(md_path, docx_path):
    """Convert markdown file to docx with proper formatting."""
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    doc = Document()

    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    style = doc.styles['Normal']
    font = style.font
    font.name = 'SimSun'
    font.size = Pt(11)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')

    software_name = "极端气候下粮食生产风险核心因子识别与韧性提升可视化决策系统"

    lines = content.split('\n')
    i = 0
    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []

    while i < len(lines):
        line = lines[i]

        if line.strip().startswith('```'):
            if in_code_block:
                if code_lines:
                    code_text = '\n'.join(code_lines)
                    p = doc.add_paragraph()
                    p.paragraph_format.space_after = Pt(4)
                    p.paragraph_format.space_before = Pt(4)
                    p.paragraph_format.left_indent = Cm(1)
                    run = p.add_run(code_text)
                    run.font.name = 'Consolas'
                    run.font.size = Pt(8)
                    run.font.color.rgb = RGBColor(64, 64, 64)
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        if line.strip().startswith('|') and line.strip().endswith('|'):
            if not in_table:
                in_table = True
                table_lines = []
            table_lines.append(line)
            i += 1
            continue
        elif in_table:
            add_table_from_markdown(doc, '\n'.join(table_lines))
            table_lines = []
            in_table = False

        if line.startswith('# ') and not line.startswith('## '):
            add_heading_styled(doc, line[2:].strip(), level=1)
        elif line.startswith('## '):
            add_heading_styled(doc, line[3:].strip(), level=2)
        elif line.startswith('### '):
            add_heading_styled(doc, line[4:].strip(), level=3)
        elif line.startswith('#### '):
            add_heading_styled(doc, line[5:].strip(), level=4)
        elif line.startswith('---'):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            pPr = p._p.get_or_add_pPr()
            pBdr = OxmlElement('w:pBdr')
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '6')
            bottom.set(qn('w:space'), '1')
            bottom.set(qn('w:color'), '999999')
            pBdr.append(bottom)
            pPr.append(pBdr)
        elif line.strip().startswith('- ') or line.strip().startswith('* '):
            text = re.sub(r'^[\s]*[-*]\s+', '', line)
            text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
            p = doc.add_paragraph(style='List Bullet')
            p.clear()
            run = p.add_run(text)
            run.font.size = Pt(10.5)
            run.font.name = 'SimSun'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')
            p.paragraph_format.space_after = Pt(3)
        elif line.strip().startswith('> '):
            text = line.strip()[2:]
            text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.5)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(text)
            run.font.size = Pt(10)
            run.font.name = 'SimSun'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')
            run.font.italic = True
            run.font.color.rgb = RGBColor(96, 96, 96)
        elif line.strip():
            text = line.strip()
            text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
            text = re.sub(r'`([^`]+)`', r'\1', text)
            text = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', text)

            if text:
                add_paragraph(doc, text)
        else:
            if not in_table:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(0)

        i += 1

    if in_table and table_lines:
        add_table_from_markdown(doc, '\n'.join(table_lines))

    add_header(doc, software_name)
    add_page_number(doc)

    doc.save(docx_path)
    print(f"Saved: {docx_path}")
    print(f"File size: {os.path.getsize(docx_path)} bytes")

if __name__ == '__main__':
    md_path = os.path.join(os.path.dirname(__file__), '草稿', '软件设计说明书.md')
    docx_path = os.path.join(os.path.dirname(__file__), '正式资料',
                             '极端气候下粮食生产风险核心因子识别与韧性提升可视化决策系统_软件设计说明书.docx')
    process_markdown_to_docx(md_path, docx_path)
